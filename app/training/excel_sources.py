"""Read workbook data without executing formulae or trusting worksheet instructions."""
import hashlib
import json
import re
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from zipfile import ZipFile, BadZipFile
from xml.etree.ElementTree import ParseError
from openpyxl.utils.exceptions import InvalidFileException

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel, WINDOWS_EPOCH


SOURCE_ERRORS = (ValueError, OSError, KeyError, BadZipFile, ParseError, InvalidFileException)

def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def normalized(value):
    value = text(value).replace('º', '').replace('ª', '').replace('°', '')
    value = unicodedata.normalize('NFKD', value).casefold()
    return re.sub(r'[^a-z0-9]', '', value)


def decimal(value):
    try:
        result = Decimal(text(value).replace(',', '.'))
    except InvalidOperation:
        raise ValueError('Número inválido') from None
    if not result.is_finite():
        raise ValueError('Número não finito')
    return result


def integer(value):
    result = decimal(value)
    if result != result.to_integral_value():
        raise ValueError('Número inteiro obrigatório')
    return int(result)


def parse_date(value, epoch=WINDOWS_EPOCH):
    if value is None or value == '':
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    for pattern in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.strptime(text(value), pattern).date()
        except ValueError:
            pass
    try:
        serial = decimal(value)
        if serial <= 0 or serial != serial.to_integral_value():
            raise ValueError('Número de série de data inválido')
        converted = from_excel(float(serial), epoch=epoch)
        if isinstance(converted, datetime):
            return converted.date()
    except (ValueError, OverflowError):
        pass
    raise ValueError('Data inválida')


def source_json(headers, row):
    return {text(key): value.isoformat() if isinstance(value, (date, datetime)) else value
            for key, value in zip(headers, row) if key is not None}


EMPLOYEE_FIELDS = {
    'number': ['Nº Mecanográfico', 'N.º Mecan', 'numero_mecanografico'],
    'name': ['Nome completo', 'Nome Oficial EPAL', 'Nome'],
    'gender': ['Género', 'Sexo'],
    'birth_date': ['Data de nascimento', 'Data de Nascimeto', 'Nascimento'],
    'identity_number': ['Nº BI', 'Nº do BI', 'BI'],
    'social_security_number': ['Nº INSS'],
    'marital_status': ['Estado civil', 'Estado'],
    'phone': ['Telemóvel', 'Nº Telemóvel', 'N.º de Telefoe', 'Telefone'],
    'email': ['E-mail', 'Email'],
    'admission_date': ['Data de admissão', 'Admissão'],
    'category': ['Categoria profissional', 'Categoria'],
    'function': ['Função exercida', 'Função'],
    'direction': ['Direcção e Gabinete', 'Direcção / Gabinete', 'Direcção'],
    'area': ['Área alocada', 'Área'],
    'company': ['Empresa'],
    'cost_center': ['Centro de Custo'],
    'registration_origin': ['Origem do cadastro'],
    'review_status': ['Estado de revisão'],
    'registration_validation': ['Validação Cadastro EPAL'],
    'source_row': ['Linha fonte EPAL'],
}
TRAINING_FIELDS = {
    'number': EMPLOYEE_FIELDS['number'],
    'participation_id': ['ID Participação'],
    'title': ['Acção de Formação', 'Nome Formação'],
    'source_course_code': ['Código Formação'],
    'area': ['Área/Categoria Formação', 'Área/Categoria'],
    'trainer': ['Formador(a)'], 'provider': ['Entidade Formadora'],
    'location': ['Local da Formação', 'Local de formação'],
    'start': ['Data de Início'], 'end': ['Data de Fim'],
    'days': ['Duração (Dias)'], 'hours_per_day': ['Horas/Dia', 'Nº Hora por dia'],
    'attendance': ['Presenças', 'Nº Presenças'],
    'attendance_confirmed': ['Presenças verificadas'],
    'grade': ['Nota da Avaliação Final', 'Média Final'],
    'result': ['Resultado Final'],
    'status': ['Situação do Formando na Formação', 'Status-Curso'],
    'plan_status': ['Status do Plano'],
    'certificate_reference': ['Nº/Ref. Certificado'],
    'certified': ['Certificado Entregue'],
    'certificate_delivery_date': ['Data de Entrega do Certificado'],
    'certificate_declaration': ['Estado certificado — declaração'],
    'certificate_declaration_source': ['Origem da confirmação'],
    'observations': ['Observações'],
    'historical_name': ['Nome Histórico Pilar 3B'],
    'identity_quality': ['Qualidade da Identidade'],
    'matching_method': ['Método de Correspondência'],
    'source_row': ['Linha Fonte Pilar 3B'],
}
IDENTIFIERS = {'number', 'identity_number', 'social_security_number', 'phone',
               'source_course_code', 'participation_id', 'certificate_reference'}
DATES = {'birth_date', 'admission_date', 'start', 'end', 'certificate_delivery_date'}
NUMBERS = {'source_row', 'days', 'attendance'}
DECIMALS = {'hours_per_day', 'grade'}


def load_source(source):
    """Bound expansion before openpyxl opens the file. No external links or VBA."""
    try:
        with ZipFile(source) as archive:
            if sum(item.file_size for item in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Excel demasiado grande após descompressão (máximo 100 MB).')
    except BadZipFile:
        raise ValueError('Ficheiro Excel inválido. Utilize .xlsx ou .xlsm.') from None
    if hasattr(source, 'seek'):
        source.seek(0)
    return load_workbook(source, read_only=True, data_only=False, keep_links=False)


def issue(report, severity, sheet, row, reason, **extra):
    report['issues'].append(dict(severity=severity, sheet=sheet, row=row, reason=reason, **extra))


def read_sheet(sheet, fields, required, epoch, report):
    """Return canonical values and original cells, with explicit row provenance."""
    if sheet.max_column and sheet.max_column > 200:
        issue(report, 'error', sheet.title, None, 'Limite de 200 colunas excedido.')
        return []
    aliases = {key: [normalized(alias) for alias in names] for key, names in fields.items()}
    positions = None
    result = []
    for row_number, cells in enumerate(sheet.iter_rows(max_col=200), 1):
        if row_number > 50000:
            issue(report, 'error', sheet.title, row_number, 'Limite de 50000 linhas excedido.')
            break
        values = [cell.value for cell in cells]
        if positions is None:
            heads = [normalized(value) for value in values]
            found = {key: [i for alias in names for i, head in enumerate(heads) if head == alias]
                     for key, names in aliases.items()}
            found = {key: list(dict.fromkeys(indices)) for key, indices in found.items()}
            if all(found[key] for key in required):
                positions, headers = found, values
                report['sheets'].append({'sheet': sheet.title, 'header_row': row_number})
            elif row_number >= 30:
                break
            continue
        if not any(value is not None and value != '' for value in values):
            continue
        raw = source_json(headers, values)
        mapped = {}
        for key, indices in positions.items():
            candidates = []
            for index in indices:
                if index >= len(cells):
                    continue
                cell = cells[index]
                value = cell.value
                if value is None or value == '':
                    continue
                try:
                    if cell.data_type == 'f':
                        raise ValueError('Fórmula numa coluna importada; forneça o valor confirmado')
                    if key in DATES:
                        value = parse_date(value, epoch)
                    elif key in NUMBERS:
                        value = integer(value)
                    elif key in DECIMALS:
                        value = decimal(value)
                    elif key in IDENTIFIERS and isinstance(value, (int, float)):
                        value = str(integer(value))
                        if re.fullmatch(r'0+', cell.number_format or ''):
                            value = value.zfill(len(cell.number_format))
                        else:
                            issue(report, 'warning', sheet.title, row_number,
                                  'Identificador numérico sem máscara: zeros já perdidos na origem não podem ser recuperados.', field=key)
                    else:
                        value = text(value)
                    if value not in candidates:
                        candidates.append(value)
                except ValueError as error:
                    issue(report, 'error', sheet.title, row_number, str(error), field=key)
            if len(candidates) > 1:
                issue(report, 'error', sheet.title, row_number,
                      'Colunas alternativas discordam na mesma linha.', field=key,
                      values=[str(value) for value in candidates])
            if candidates:
                mapped[key] = candidates[0]
        result.append({'sheet': sheet.title, 'row': row_number, 'values': mapped, 'raw': raw})
    if positions is None:
        issue(report, 'error', sheet.title, None, 'Cabeçalho reconhecível não encontrado nas primeiras 30 linhas.')
    return result


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()
