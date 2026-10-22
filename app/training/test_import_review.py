import hashlib
import json
import tempfile
from datetime import date
from io import BytesIO, StringIO
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.core.management import call_command
from django.core.management.base import CommandError
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, Client
from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904

from .excel_sources import parse_date
from .integrated_import import apply_reviewed, build_plan
from .models import AuditEvent, Course, Employee, Training


def sample():
    book = Workbook()
    employees = book.active
    employees.title = 'FUNCIONARIOS'
    employees.append(['Nome completo', 'Nº Mecanográfico', 'Nº BI', 'Nº INSS', 'Função exercida',
                      'Data de nascimento', 'Género', 'Estado civil', 'Telemóvel', 'E-mail',
                      'Data de admissão', 'Categoria profissional', 'Direcção e Gabinete',
                      'Área alocada', 'Origem do cadastro', 'Estado de revisão', 'Linha fonte EPAL'])
    employees.append(['Técnica Teste', '001', '000123LA', '001234', 'Operadora', '29485', 'F',
                      'Solteira', '900000000', 'teste@example.com', '40857', 'Técnica', 'Direcção',
                      'Área A', 'Cadastro', 'Apto', '20'])
    secondary = book.create_sheet('FUNCIONARIOS_EPAL')
    secondary.append(['N.º Mecan', 'Nome', 'Função', 'Validação Cadastro EPAL'])
    secondary.append(['001', 'Técnica Teste', 'Operadora', 'Definitivo'])
    participation = book.create_sheet('PARTICIPACOES')
    participation.append(['ID Participação', 'Nº Mecanográfico', 'Acção de Formação', 'Código Formação',
                          'Área/Categoria Formação', 'Data de Início', 'Data de Fim', 'Duração (Dias)',
                          'Horas/Dia', 'Presenças', 'Situação do Formando na Formação', 'Empresa',
                          'Centro de Custo', 'Estado certificado — declaração', 'Origem da confirmação'])
    participation.append(['PF-001', '001', 'Segurança', 'SEG', 'Técnica', '45658', '45659', '2',
                          '7', '2', 'Concluída', 'EPAL', 'CC01', 'Emitido e entregue', 'Declaração do responsável'])
    copy = book.create_sheet('BASE_MESTRA_FORMACAO')
    for row in participation.values:
        copy.append(row)
    return book


def encoded(book):
    data = BytesIO()
    book.save(data)
    return data.getvalue()


class ImportReviewTests(TestCase):
    def setUp(self):
        self.book = sample()
        self.sha = 'sample-sha'
        self.course = Course.objects.create(code='SEG', title='Segurança')

    def plan(self):
        return build_plan(self.book, self.sha)

    def apply(self):
        plan = self.plan()
        self.assertEqual(plan.report['errors'], 0, plan.report['issues'])
        return apply_reviewed(self.book, self.sha, plan.report['review_token'])

    def test_all_employee_fields_reordered_headers_and_identifiers(self):
        before = (Employee.objects.count(), Training.objects.count(), AuditEvent.objects.count())
        plan = self.plan()
        self.assertEqual((Employee.objects.count(), Training.objects.count(), AuditEvent.objects.count()), before)
        self.assertTrue(plan.report['can_apply'], plan.report['issues'])
        self.assertEqual(plan.report['source_counts'], {'employees': 1, 'participations': 1})
        self.assertEqual(plan.report['duplicates_reconciled'], 2)
        obj = plan.employees[0][0]
        expected = {'number': '001', 'name': 'Técnica Teste', 'identity_number': '000123LA',
                    'social_security_number': '001234', 'function': 'Operadora', 'gender': 'F',
                    'marital_status': 'Solteira', 'phone': '900000000', 'email': 'teste@example.com',
                    'category': 'Técnica', 'direction': 'Direcção', 'area': 'Área A',
                    'registration_origin': 'Cadastro', 'review_status': 'Apto', 'source_row': 20,
                    'company': 'EPAL', 'cost_center': 'CC01', 'registration_validation': 'Definitivo'}
        for field, value in expected.items():
            self.assertEqual(getattr(obj, field), value)
        self.assertEqual(obj.birth_date, parse_date(29485))
        self.assertEqual(obj.admission_date, parse_date(40857))

    def test_idempotence_catalogue_and_declarations(self):
        self.apply()
        self.assertEqual(Employee.objects.count(), 1)
        self.assertEqual(Training.objects.count(), 1)
        training = Training.objects.get()
        self.assertEqual(training.course_id, self.course.pk)
        self.assertFalse(training.attendance_confirmed)
        self.assertIsNone(training.certified)
        self.assertIsNone(training.certificate_issued)
        self.assertEqual(training.certificate_declaration, 'Emitido e entregue')
        self.assertEqual(len(training.source_data['_import']['sources']), 2)
        report = self.apply()
        self.assertEqual(report['trainings_created'], 0)
        self.assertEqual(report['employees_updated'], 0)
        self.assertEqual(report['totals']['completed_hours_declared'], '14')
        self.assertEqual(report['totals']['certificate_declarations'], 1)
        self.assertEqual(Course.objects.count(), 1)

    def test_copies_and_repeated_ids_block_all_writes(self):
        self.book['BASE_MESTRA_FORMACAO']['I2'] = 6
        plan = self.plan()
        self.assertFalse(plan.report['can_apply'])
        with self.assertRaises(ValueError):
            apply_reviewed(self.book, self.sha, plan.report['review_token'])
        self.assertFalse(Employee.objects.exists())
        self.assertFalse(Training.objects.exists())
        self.book['BASE_MESTRA_FORMACAO']['I2'] = '7'
        self.book['PARTICIPACOES'].append(list(self.book['PARTICIPACOES'].values)[1])
        self.assertFalse(self.plan().report['can_apply'])

    def test_existing_nonempty_values_protected_and_blank_cells_do_not_erase(self):
        self.apply()
        employee = Employee.objects.get()
        employee.phone = '999999999'
        employee.save()
        self.assertFalse(self.plan().report['can_apply'])
        self.book['FUNCIONARIOS']['I2'] = None
        plan = self.plan()
        self.assertTrue(plan.report['can_apply'], plan.report['issues'])
        self.assertEqual(plan.employees[0][0].phone, '999999999')
        self.apply()
        employee.refresh_from_db()
        self.assertEqual(employee.phone, '999999999')

    def test_changed_database_or_file_requires_new_review(self):
        token = self.plan().report['review_token']
        self.course.area = 'Alterada'
        self.course.save()
        with self.assertRaisesMessage(ValueError, 'mudou'):
            apply_reviewed(self.book, self.sha, token)
        token = self.plan().report['review_token']
        with self.assertRaisesMessage(ValueError, 'mudou'):
            apply_reviewed(self.book, 'other-file-hash', token)
        self.assertFalse(Employee.objects.exists())

    def test_atomic_rollback_on_late_write_error(self):
        token = self.plan().report['review_token']
        with patch('training.integrated_import.AuditEvent.objects.create', side_effect=RuntimeError('late error')):
            with self.assertRaises(RuntimeError):
                apply_reviewed(self.book, self.sha, token)
        self.assertFalse(Employee.objects.exists())
        self.assertFalse(Training.objects.exists())

    def test_numeric_identifiers_with_zero_mask_and_1904_dates(self):
        self.book.epoch = CALENDAR_MAC_1904
        for sheet, cell in [('FUNCIONARIOS', 'B2'), ('FUNCIONARIOS_EPAL', 'A2'),
                            ('PARTICIPACOES', 'B2'), ('BASE_MESTRA_FORMACAO', 'B2')]:
            self.book[sheet][cell] = 1
            self.book[sheet][cell].number_format = '000'
        plan = self.plan()
        self.assertEqual(plan.employees[0][0].number, '001')
        self.assertEqual(plan.employees[0][0].birth_date, parse_date('29485', CALENDAR_MAC_1904))
        self.assertEqual(parse_date('2025-01-01'), date(2025, 1, 1))
        for value in ('NaN', '-1', 'nonsense', 'Infinity'):
            with self.assertRaises(ValueError):
                parse_date(value)

    def test_alias_conflict_formula_and_invalid_row_block_apply(self):
        self.book['FUNCIONARIOS']['R1'] = 'Função'
        self.book['FUNCIONARIOS']['R2'] = 'Outra'
        self.assertFalse(self.plan().report['can_apply'])
        self.book['FUNCIONARIOS']['R2'] = '=1+1'
        self.assertTrue(any('Fórmula' in i['reason'] for i in self.plan().report['issues']))
        self.book['FUNCIONARIOS']['R2'] = None
        self.book['FUNCIONARIOS']['J2'] = 'not-an-email'
        plan = self.plan()
        self.assertFalse(plan.report['can_apply'])
        with self.assertRaises(ValueError):
            apply_reviewed(self.book, self.sha, plan.report['review_token'])
        self.assertFalse(Employee.objects.exists())

    def test_missing_id_and_wrong_catalogue_identity(self):
        self.book['PARTICIPACOES']['A2'] = None
        self.assertFalse(self.plan().report['can_apply'])
        self.book['PARTICIPACOES']['A2'] = 'PF-001'
        self.course.title = 'Different training'
        self.course.save()
        self.assertFalse(self.plan().report['can_apply'])

    def test_command_review_token_no_source_changes_and_safe_default(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / 'source.xlsx'
            report = Path(folder) / 'preview.json'
            content = encoded(self.book)
            path.write_bytes(content)
            call_command('import_excel', path, report=report, stdout=StringIO())
            self.assertFalse(Employee.objects.exists())
            preview = json.loads(report.read_text())
            self.assertEqual(preview['source_sha256'], hashlib.sha256(content).hexdigest())
            with self.assertRaises(CommandError):
                call_command('import_excel', path, report=report, apply=True, stdout=StringIO())
            call_command('import_excel', path, report=report, apply=True,
                         review_token=preview['review_token'], stdout=StringIO())
            self.assertEqual(Employee.objects.count(), 1)
            self.assertEqual(path.read_bytes(), content)

    def test_preview_permissions_csrf_download_and_no_write(self):
        url = '/administracao/importacao/'
        user = get_user_model().objects.create_user('staff', is_staff=True)
        self.client.force_login(user)
        self.assertEqual(self.client.get(url).status_code, 403)
        user.is_superuser = True
        user.save()
        self.assertContains(self.client.get('/admin/'), url)
        csrf_client = Client(enforce_csrf_checks=True)
        csrf_client.force_login(user)
        self.assertEqual(csrf_client.post(url, {}).status_code, 403)
        self.book['FUNCIONARIOS']['A2'] = '<script>alert(1)</script>'
        self.book['FUNCIONARIOS_EPAL']['B2'] = '<script>alert(1)</script>'
        data = encoded(self.book)
        response = self.client.post(url, {'workbook': SimpleUploadedFile('test.xlsx', data)})
        self.assertContains(response, '&lt;script&gt;')
        self.assertNotContains(response, '<script>alert(1)</script>')
        self.assertIn('no-store', response['Cache-Control'])
        response = self.client.post(url, {'workbook': SimpleUploadedFile('test.xlsx', data), 'output': 'json'})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()['mode'], 'dry-run')
        self.assertFalse(Employee.objects.exists())
        self.assertFalse(Training.objects.exists())
        self.assertFalse(AuditEvent.objects.exists())
