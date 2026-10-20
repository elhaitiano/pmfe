"""Reviewed catalogue-only import. No history or group models are imported here."""
import hashlib
import json
from collections import Counter
from decimal import Decimal, InvalidOperation

from django.core.exceptions import ValidationError
from django.core.management.base import CommandError
from django.db import IntegrityError, OperationalError, connections, transaction

from .catalogue_diagnostic import normalise
from .models import AuditEvent, Course


SOURCE_SYSTEM = 'epal-proagua-catalogo'
PLAN_VERSION = 1
HEADERS = ['ID Acção', 'Acção de formação', 'Código interno', 'Área',
           'Dias padrão declarados', 'Horas padrão disponíveis', 'Origem', 'Estado designação']
CATALOGUE_FIELDS = ('id', 'title', 'code', 'area', 'days', 'hours', 'origin', 'designation_state')


def digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False,
                                     separators=(',', ':'), default=str).encode('utf-8')).hexdigest()


def read_catalogue(path):
    """Read only CATALOGO, never save the workbook or inspect staff worksheets."""
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        if 'CATALOGO' not in workbook.sheetnames:
            raise CommandError('Folha CATALOGO não encontrada.')
        sheet = workbook['CATALOGO']
        header = next(sheet.iter_rows(min_row=3, max_row=3, max_col=8, values_only=True))
        if list(header) != HEADERS:
            raise CommandError('Cabeçalhos de CATALOGO diferentes da estrutura revista.')
        rows = []
        for values in sheet.iter_rows(min_row=4, max_col=8, values_only=True):
            if not any(v is not None for v in values):
                continue
            row = {}
            for key, value in zip(CATALOGUE_FIELDS, values):
                row[key] = (None if value is None else str(value)) if key in ('days', 'hours') else str(value or '').strip()
            rows.append(row)
        return rows
    finally:
        workbook.close()


def snapshot(database):
    # Only catalogue metadata. This also detects additions/deletions and manual edits.
    return list(Course.objects.using(database).order_by('pk').values())


def database_identity(database):
    config = connections[database].settings_dict
    return {'alias': database, 'engine': config['ENGINE'], 'name': str(config['NAME'])}


def make_plan(rows, current, destination, reviewed_decisions=None):
    ids = Counter(r['id'] for r in rows)
    titles = Counter(normalise(r['title']) for r in rows)
    codes = Counter(r['code'].casefold() for r in rows if r['code'])
    identities = {(c['source_system'], c['source_action_id']): c for c in current
                  if c['source_system'] and c['source_action_id']}
    operations, conflicts = [], []
    for row in rows:
        problems = []
        if not row['id'] or len(row['id']) > 100 or ids[row['id']] != 1:
            problems.append('ID de origem ausente, inválido ou repetido.')
        if not row['title'] or len(row['title']) > 1000 or titles[normalise(row['title'])] != 1:
            problems.append('Designação ausente, inválida ou repetida após normalização.')
        if len(row['code']) > 100 or (row['code'] and codes[row['code'].casefold()] != 1):
            problems.append('Código de origem inválido ou repetido.')
        if len(row['area']) > 200:
            problems.append('Área excede o limite do catálogo.')
        days = None
        try:
            if row['days'] is not None:
                number = Decimal(row['days'])
                if not number.is_finite() or number != number.to_integral_value() or not 1 <= number <= 366:
                    raise ValueError
                days = int(number)
            if row['hours'] is not None:
                hours = Decimal(row['hours'])
                if not hours.is_finite() or hours <= 0:
                    raise ValueError
        except (ValueError, InvalidOperation):
            problems.append('Duração de origem inválida.')
        existing = identities.get((SOURCE_SYSTEM, row['id']))
        fingerprint = digest(row)
        for course in current:
            if existing and course['id'] == existing['id']:
                continue
            known_titles = {normalise(t) for t in [course['title'], *course['aliases']]}
            if normalise(row['title']) in known_titles:
                problems.append('Título ou alias já pertence a outra acção; associação manual necessária.')
            known_codes = {str(course[f]).casefold() for f in ('internal_code', 'source_internal_code') if course[f]}
            if row['code'] and row['code'].casefold() in known_codes:
                problems.append('Código já pertence a outra acção; revisão necessária.')
        if existing and existing['source_digest'] != fingerprint:
            problems.append('A origem mudou; não substituir a versão já integrada automaticamente.')
        status = Course.ValidationState.PENDING
        notes = ['Importação do catálogo não constitui validação.',
                 'Horas padrão disponíveis preservadas na origem; confirmar unidade antes de definir horas por dia.',
                 'Aliases históricos não aprovados automaticamente.']
        if not row['code']:
            notes.append('Código interno ausente; não atribuído.')
        if not row['area']:
            notes.append('Área por confirmar.')
        if row['designation_state'] == 'Designação por confirmar':
            status = Course.ValidationState.BLOCKED
            notes.append('Designação da origem por confirmar.')
        if row['code'].upper() in ('WSFR', 'WSFT') or row['id'] == 'AC-2026-051':
            status = Course.ValidationState.BLOCKED
            notes.append('Conflito WSFR/WSFT: decisão documental obrigatória antes de validar.')
        values = dict(source_system=SOURCE_SYSTEM, source_action_id=row['id'],
                      source_internal_code=row['code'], source_digest=fingerprint,
                      source_metadata=dict(row), internal_code='', title=row['title'],
                      area=row['area'], contents='', default_days=days, default_hours=None,
                      validation_status=status, validation_notes='\n'.join(notes),
                      aliases=[], active=True)
        operations.append({'source_action_id': row['id'],
                           'operation': 'conflict' if problems else 'keep' if existing else 'create',
                           'existing_id': existing['id'] if existing else None,
                           'values': values, 'problems': problems})
        conflicts.extend({'source_action_id': row['id'], 'reason': p} for p in problems)
    plan = dict(version=PLAN_VERSION, source_system=SOURCE_SYSTEM,
                destination_database=destination,
                source_digest=digest(rows), destination_digest=digest(current),
                operations=operations, conflicts=conflicts,
                summary=dict(records=len(rows), creates=sum(o['operation'] == 'create' for o in operations),
                             keeps=sum(o['operation'] == 'keep' for o in operations),
                             conflicts=len(conflicts)))
    if reviewed_decisions is not None:
        review_codes(plan, reviewed_decisions, current)
    plan['plan_digest'] = digest(plan)
    return plan


def review_codes(plan, decisions, current):
    """Apply explicit code decisions, bound to unchanged source metadata.

    This approves a code, never a title, alias, duration or complete action.
    Other operations remain byte-for-byte equivalent in the plan payload.
    """
    if not isinstance(decisions, list):
        raise CommandError('As decisões devem ser uma lista de confirmações explícitas.')
    required = {'source_action_id', 'source_digest', 'source_internal_code',
                'internal_code', 'approval_note'}
    by_id = {o['source_action_id']: o for o in plan['operations']}
    seen = set()
    for decision in decisions:
        if (not isinstance(decision, dict) or set(decision) != required or
                any(not isinstance(v, str) or not v.strip() for v in decision.values())):
            raise CommandError('Decisão inválida: indique ID, hash da origem, códigos e confirmação documentada.')
        identifier = decision['source_action_id']
        if identifier in seen or identifier not in by_id:
            raise CommandError('Decisão repetida ou sem correspondência na origem.')
        seen.add(identifier)
        operation = by_id[identifier]
        values = operation['values']
        if (operation['operation'] not in ('create', 'keep') or
                values['source_digest'] != decision['source_digest'] or
                values['source_internal_code'] != decision['source_internal_code']):
            raise CommandError('A decisão não corresponde à origem revista intacta e sem conflitos.')
        code = decision['internal_code']
        if code != code.strip() or len(code) > 100:
            raise CommandError('Código aprovado inválido.')
        if operation['operation'] == 'keep':
            existing = next(c for c in current if c['id'] == operation['existing_id'])
            if existing['internal_code'] != code:
                raise CommandError('A decisão alteraria uma acção existente; revisão separada necessária.')
        for other in plan['operations']:
            if other is operation:
                continue
            other_values = other['values']
            if code.casefold() in {other_values['source_internal_code'].casefold(),
                                  other_values['internal_code'].casefold()}:
                raise CommandError('Código aprovado colide com outra acção do plano.')
        # The existing destination codes were checked against source codes above;
        # check newly approved codes too, including codes outside the source sheet.
        destination_codes = {str(c[f]).casefold() for c in current
                             if c['id'] != operation['existing_id']
                             for f in ('internal_code', 'source_internal_code') if c[f]}
        if code.casefold() in destination_codes:
            raise CommandError('Código aprovado já existe no destino.')
        values['internal_code'] = code
        notes = values['validation_notes'].splitlines()
        resolved = (identifier == 'AC-2026-051' and
                    values['source_internal_code'] == 'WSFR' and code == 'WSFT')
        if resolved:
            notes = [n for n in notes if not n.startswith('Conflito WSFR/WSFT:')]
            if values['source_metadata']['designation_state'] != 'Designação por confirmar':
                values['validation_status'] = Course.ValidationState.PENDING
        notes.append(f'Código original preservado: {values["source_internal_code"]}. Código aprovado: {code}.')
        notes.append('Confirmação expressa do responsável: ' + decision['approval_note'])
        notes.append('Decisão limitada ao código; restantes pendências de validação mantidas.')
        values['validation_notes'] = '\n'.join(notes)
    plan['reviewed_decisions'] = decisions


def simulate(rows, database='default', reviewed_decisions=None):
    from .management.commands.diagnose_catalogue import select_only
    with connections[database].execute_wrapper(select_only):
        return make_plan(rows, snapshot(database), database_identity(database), reviewed_decisions)


def apply_plan(rows, plan, approval, actor, database='default'):
    """Explicit approval required. Creates catalogue entries only; never updates history."""
    content = {k: v for k, v in plan.items() if k != 'plan_digest'}
    if (not actor or not actor.strip() or not approval or
            approval != plan.get('plan_digest') or digest(content) != approval):
        raise CommandError('É obrigatório aprovar explicitamente o hash de um plano de simulação íntegro.')
    if plan.get('conflicts'):
        raise CommandError('Plano com conflitos: nenhuma aplicação permitida.')
    try:
        with transaction.atomic(using=database):
            # First write reserves SQLite's writer lock BEFORE checking the snapshot.
            # A rejected plan rolls this event back together with the entire batch.
            event = AuditEvent.objects.using(database).create(
                actor=actor.strip()[:150], action='CATALOGUE_IMPORT', detail='Pending validation')
            courses = Course.objects.using(database).select_for_update().order_by('pk')
            current = list(courses.values())
            expected = make_plan(rows, current, database_identity(database), plan.get('reviewed_decisions'))
            if expected['plan_digest'] != approval:
                raise CommandError('Origem, catálogo ou plano alterado. Execute e aprove uma nova simulação.')
            created = []
            for operation in expected['operations']:
                if operation['operation'] != 'create':
                    continue
                course = Course(**operation['values'])
                course._state.db = database
                # Title identity is generated by save; validate its other fields first.
                course.key = normalise(course.title)
                course.full_clean()
                course.save(using=database, force_insert=True)
                created.append(course.source_action_id)
            event.detail = json.dumps({'plan_digest': approval, 'source_digest': plan['source_digest'],
                                       'created': created, 'kept': plan['summary']['keeps']}, ensure_ascii=False)
            event.save(using=database, update_fields=['detail'])
            return {'created': len(created), 'kept': plan['summary']['keeps'], 'audit_id': event.pk}
    except (IntegrityError, OperationalError, ValidationError) as exc:
        raise CommandError('Aplicação anulada integralmente; reveja dados ou concorrência e simule novamente.') from exc
