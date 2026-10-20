"""Synthetic reports, equivalent grades and individual correction security."""
import hashlib,io,json,time
from datetime import date
from decimal import Decimal
from unittest.mock import patch
from django.contrib.auth import get_user_model
from django.contrib.auth.models import Permission
from django.core.exceptions import ValidationError,PermissionDenied
from django.db import connection
from django.db.models import TextField
from django.db.models.functions import Cast
from django.test import TestCase,Client,override_settings
from django.urls import reverse
from .models import Employee,Training,Course,AuditEvent,TrainingGroup,Enrollment,GroupSession
from . import consultations as c,historical as h,evaluation_corrections as corrections
from .approved_evaluations import binding,ITEM
from .source_achievement import attach
from .evaluations import percentage_to_grade,result

@override_settings(PMF_HISTORICAL_ENABLED=True,PMF_DOCUMENTS_ENABLED=True)
class ConsultationTests(TestCase):
    def setUp(self):
        self.admin=get_user_model().objects.create_superuser('reports-admin',password='synthetic')
        self.reader=get_user_model().objects.create_user('reports-reader',password='synthetic')
        self.reader.user_permissions.add(*Permission.objects.filter(codename__in=['view_employee','view_training']))
        self.noaccess=get_user_model().objects.create_user('reports-denied')
        self.people=[Employee.objects.create(number='SYN-REP-'+str(i),name='Synthetic person '+str(i),direction='Synthetic direction',birth_date=date(1990,1,1),source_data={'Validação Cadastro EPAL':'Definitivo — synthetic','Data de Nascimeto':'1990-01-01'}) for i in range(3)]
        self.course=Course.objects.create(title='Synthetic report training',source_system='epal-proagua-catalogo',source_action_id='SYN-REPORT')
        self.rows=[self.training(self.people[0],'SYN-REP-1'),self.training(self.people[0],'SYN-REP-2'),self.training(self.people[1],'SYN-REP-3'),self.training(self.people[2],'SYN-REP-4'),self.training(self.people[0],'pilar3b:PF-00285')]
        self.supplement(self.rows[0],'percentage','80','80')
        self.supplement(self.rows[2],'qualitative','Bom',None)
        self.manifest={'excluded_keys':list(c.EXCLUDED_KEYS),'editions':{'SYN-E':{'id':'SYN-E','action_id':'SYN-REPORT','display_id':'HR-SYN','title':'Synthetic report training','documented_group':'','start':'2026-01-01','end':'2026-01-02','source_origin':'Synthetic','location':'','trainer':''}},'bindings':{}}
        for t in self.rows[:-1]:
            obj=Training.objects.annotate(source_json_text=Cast('source_data',TextField())).get(pk=t.pk)
            self.manifest['bindings'][t.source_key]={'edition':'SYN-E','phase':'SYN-P','action_id':'SYN-REPORT','level':'probable','signature':h.signature(obj),'employee_number_sha256':hashlib.sha256(t.employee.number.encode()).hexdigest()}
        mock=patch('training.historical.manifest',return_value=self.manifest);mock.start();self.addCleanup(mock.stop)
    def training(self,person,key):
        return Training.objects.create(employee=person,title='Synthetic report training',source_key=key,start=date(2026,1,1),end=date(2026,1,2),days=2,hours_per_day=7,attendance=2,attendance_confirmed=True,status='Concluído',source_data={'original':'preserved'})
    def supplement(self,t,unit,raw,percentage):
        with connection.cursor() as cursor:cursor.execute('select source_data from training_training where id=%s',[t.pk]);source=cursor.fetchone()[0]
        AuditEvent.objects.create(actor='synthetic',action=ITEM,detail=json.dumps({'training_id':t.pk,'source_key':t.source_key,'signature':binding(t,source),'unit':unit,'raw':raw,'percentage':percentage,'approved':True,'source_sheet':'PARTICIPACOES','evidence_cells':[{'format':'0%','ref':'Base de dados:10','cell':'AE10','type':'n'}]}))
    def data(self,**changes):
        d={'unit':'percentage','value':'75','reason':'Correcção documental sintética','reference':'Documento sintético 1','final_achievement':True};d.update(changes);return d
    def token(self,t=None,**changes):
        t=t or Training.objects.get(pk=self.rows[0].pk)
        return corrections.preview(self.admin,t,self.data(**changes))[0]
    def login(self,user=None):self.client.force_login(user or self.reader)
    def test_decimal_conversion_and_rounding(self):
        for value,expected in [('80','16.00'),('0','0.00'),('100','20.00'),('83.33','16.67')]:self.assertEqual(percentage_to_grade(value),Decimal(expected))
        for value in ['NaN','Infinity','-1','101','unknown']:self.assertIsNone(percentage_to_grade(value))
    def test_approved_percentage_converted_without_write(self):
        count=AuditEvent.objects.count();rows=c.select({})
        t=next(t for t in rows if t.pk==self.rows[0].pk)
        self.assertEqual(t.evaluation['numeric'],16);self.assertTrue(t.evaluation['calculated']);self.assertEqual(t.evaluation['percentage'],80)
        self.assertIsNone(Training.objects.get(pk=t.pk).grade);self.assertEqual(AuditEvent.objects.count(),count)
    def test_unapproved_or_nonfinal_percentage_never_converted(self):
        t=self.rows[0];t.source_achievement={'unit':'percentage','percentage':Decimal(80),'raw':'80','approved':False,'source_sheet':'PARTICIPACOES','evidence_cells':[{'format':'0%'}]}
        self.assertIsNone(result(t)['numeric'])
        t.source_achievement['approved']=True;t.source_achievement['source_sheet']='Satisfaction'
        self.assertIsNone(result(t)['numeric'])
        t.source_achievement['source_sheet']='PARTICIPACOES';t.status='Planeado'
        self.assertIsNone(result(t)['numeric'])
        t.source_achievement={'unit':'source_scale_unconfirmed','percentage':None,'raw':'70'}
        self.assertEqual(result(t)['original'],'70');self.assertEqual(result(t)['label'],'70 (escala por confirmar)');self.assertFalse(result(t)['available'])
    def test_original_grade_has_priority(self):
        t=self.rows[0];attach([t]);t.grade=Decimal(13)
        self.assertEqual(result(t)['numeric'],13);self.assertFalse(result(t)['calculated'])
    def test_attendance_column_is_never_converted_as_final_grade(self):
        t=self.rows[0];attach([t]);t.source_achievement['evidence_cells']=[{'format':'0%','ref':'Base de dados:10','cell':'AB10','type':'n'}]
        self.assertIsNone(result(t)['numeric'])
    def test_qualitative_not_missing_or_numeric(self):
        t=next(t for t in c.select({}) if t.pk==self.rows[2].pk)
        self.assertTrue(t.evaluation['available']);self.assertEqual(t.evaluation['qualitative'],'Bom');self.assertIsNone(t.evaluation['numeric'])
    def test_details_preserve_duplicates_totals_do_not(self):
        rows=c.select({});s=c.summary(rows)
        self.assertEqual(len(rows),5);self.assertEqual(s['participations'],4);self.assertEqual(s['people'],3);self.assertEqual(s['duplicates'],1)
    def test_missing_workers_exclude_people_with_any_available_result(self):
        rows=c.select({'mode':'missing'})
        self.assertEqual({t.employee_id for t in rows},{self.people[2].pk})
    def test_filters_search_period_age_direction_action_and_group(self):
        rows=c.select({'q':self.people[0].number,'year':'2026','direction':'Synthetic direction','age_band':'35','action':'SYN-REPORT','edition':'SYN-E','date_from':date(2026,1,2),'date_to':date(2026,1,3)})
        self.assertEqual(len(rows),2)
        self.assertEqual(c.select({'date_from':date(2026,2,1)}),[])
    def test_unknown_certificates_not_counted_as_delivered(self):
        self.assertEqual(c.select({'mode':'issued'}),[]);self.assertEqual(c.select({'mode':'learner'}),[]);self.assertEqual(c.select({'mode':'hr'}),[])
    def test_authentication_and_read_permissions(self):
        self.assertEqual(self.client.get(reverse('consultations')).status_code,302)
        self.login(self.noaccess);self.assertEqual(self.client.get(reverse('consultations')).status_code,403)
        self.login();self.assertEqual(self.client.get(reverse('consultations')).status_code,200)
    def test_invalid_filters_export_refused(self):
        self.login();self.assertEqual(self.client.get(reverse('consultations'),{'date_from':'2026-02-01','date_to':'2026-01-01','format':'xlsx'}).status_code,400)
        self.assertEqual(self.client.get(reverse('consultations'),{'format':'unknown'}).status_code,400)
    def test_excel_real_rows_scales_and_formula_injection_protection(self):
        from openpyxl import load_workbook
        Employee.objects.filter(pk=self.people[0].pk).update(name='=HYPERLINK("unsafe")')
        self.login();response=self.client.get(reverse('consultations'),{'format':'xlsx'})
        content=b''.join(response.streaming_content);response.close();wb=load_workbook(io.BytesIO(content))
        sheet=wb['Participações'];self.assertEqual(sheet.max_row,6)
        self.assertEqual(sheet['B2'].data_type,'s');self.assertEqual(sheet['B2'].value,'=HYPERLINK("unsafe")')
        self.assertEqual(sheet['K2'].value,16);self.assertEqual(sheet['L2'].value,'Calculada');wb.close()
    def test_pdf_export_and_denominators(self):
        self.login();response=self.client.get(reverse('consultations'),{'format':'pdf'})
        content=b''.join(response.streaming_content);response.close();self.assertTrue(content.startswith(b'%PDF'))
        self.assertIn('no-store',response['Cache-Control'])
    def test_read_only_user_never_gets_correction_links(self):
        self.login();response=self.client.get(reverse('consultations'))
        self.assertNotContains(response,'Corrigir avaliação');self.assertContains(response,'calculada')
        self.assertEqual(self.client.get(reverse('evaluation_correct',args=[self.rows[0].pk])).status_code,403)
    def test_preview_does_not_write_and_works_without_javascript(self):
        self.login(self.admin);count=AuditEvent.objects.count()
        response=self.client.post(reverse('evaluation_correct',args=[self.rows[0].pk]),self.data())
        self.assertContains(response,'Pré-visualização');self.assertContains(response,'Confirmar correcção');self.assertEqual(AuditEvent.objects.count(),count)
    def test_confirmation_audits_before_after_and_preserves_original(self):
        t=self.rows[0];original=t.source_data.copy();event=corrections.confirm(self.admin,self.token(),t.pk)
        detail=json.loads(event.detail);self.assertEqual(detail['reason'],self.data()['reason']);self.assertEqual(detail['reference'],self.data()['reference'])
        t.refresh_from_db();self.assertEqual(t.version,2);self.assertEqual(t.source_data,original);self.assertIsNone(t.grade)
        attach([t]);self.assertEqual(t.evaluation['numeric'],15);self.assertEqual(t.source_achievement['original_result']['raw'],'80')
        self.assertEqual(TrainingGroup.objects.count()+Enrollment.objects.count()+GroupSession.objects.count(),0)
    def test_repeated_confirmation_is_blocked(self):
        token=self.token();corrections.confirm(self.admin,token,self.rows[0].pk);count=AuditEvent.objects.count()
        with self.assertRaises(ValidationError):corrections.confirm(self.admin,token,self.rows[0].pk)
        self.assertEqual(AuditEvent.objects.count(),count)
    def test_stale_revision_and_binding_are_blocked(self):
        token=self.token();Training.objects.filter(pk=self.rows[0].pk).update(version=2)
        with self.assertRaises(ValidationError):corrections.confirm(self.admin,token,self.rows[0].pk)
    def test_cross_user_and_cross_participation_tokens_are_blocked(self):
        token=self.token();other=get_user_model().objects.create_superuser('reports-other',password='synthetic')
        with self.assertRaises(PermissionDenied):corrections.confirm(other,token,self.rows[0].pk)
        with self.assertRaises(PermissionDenied):corrections.confirm(self.admin,token,self.rows[1].pk)
    def test_tampered_and_expired_token(self):
        token=self.token()
        with self.assertRaises(ValidationError):corrections.confirm(self.admin,token+'bad',self.rows[0].pk)
        with patch('django.core.signing.time.time',return_value=time.time()+1000):
            with self.assertRaises(ValidationError):corrections.confirm(self.admin,token,self.rows[0].pk)
    def test_rollback_if_final_audit_save_fails(self):
        token=self.token(unit='grade20',value='12');before=AuditEvent.objects.count();save=AuditEvent.save
        def failing(event,*a,**kw):
            if event.action==corrections.ACTION:raise RuntimeError('Synthetic audit failure')
            return save(event,*a,**kw)
        with patch.object(AuditEvent,'save',failing):
            with self.assertRaises(RuntimeError):corrections.confirm(self.admin,token,self.rows[0].pk)
        self.assertEqual(AuditEvent.objects.count(),before);self.assertIsNone(Training.objects.get(pk=self.rows[0].pk).grade)
    def test_csrf_and_explicit_confirmation(self):
        client=Client(enforce_csrf_checks=True);client.force_login(self.admin)
        url=reverse('evaluation_correct',args=[self.rows[0].pk]);self.assertEqual(client.post(url,self.data()).status_code,403)
        self.login(self.admin);response=self.client.post(url,{'token':self.token()});self.assertContains(response,'Confirme expressamente')
    def test_clear_result_does_not_resurrect_original_or_certificate_states(self):
        event=corrections.confirm(self.admin,self.token(unit='none',value=''),self.rows[0].pk)
        t=Training.objects.get(pk=self.rows[0].pk);attach([t]);self.assertFalse(t.evaluation['available']);self.assertIsNone(t.certified)
    def test_correction_form_requires_document_and_final_meaning(self):
        for change in [{'reference':''},{'reason':''},{'final_achievement':False},{'unit':'grade20','value':'21'},{'value':'NaN'}]:
            with self.subTest(change=change):self.assertFalse(corrections.CorrectionForm(self.data(**change)).is_valid())
    def test_old_form_cannot_bypass_supplement_or_historical_documentation(self):
        from .forms import TrainingForm
        from django.forms.models import model_to_dict
        t=self.rows[0];data=model_to_dict(t);data.update(grade='20',certified='unknown',revision=t.version)
        form=TrainingForm(data,instance=t)
        self.assertTrue(form.fields['grade'].disabled)
        self.assertTrue(form.is_valid(),form.errors);form.save();t.refresh_from_db();self.assertIsNone(t.grade)
        t.source_key='pilar3b:SYN-PROTECTED';t.save(update_fields=['source_key'])
        data=model_to_dict(t);data.update(grade='20',certified='true',certificate_reference='Unproved',certificate_delivery_date='2026-01-02')
        form=TrainingForm(data,instance=t)
        self.assertTrue(form.is_valid(),form.errors);form.save();t.refresh_from_db()
        self.assertIsNone(t.grade);self.assertIsNone(t.certified);self.assertEqual(t.certificate_reference,'')
