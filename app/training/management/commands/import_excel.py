"""Read the original workbook; never execute VBA or save the source file."""
import hashlib
import json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from django.core.exceptions import ValidationError
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from openpyxl import load_workbook
from training.models import AuditEvent, Employee, Training


def text(value):
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def parse_date(value):
    if not value:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int,float)):
        from openpyxl.utils.datetime import from_excel
        converted = from_excel(value)
        if isinstance(converted, datetime):
            return converted.date()
    for pattern in ('%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%Y-%m-%d %H:%M:%S'):
        try:
            return datetime.strptime(str(value).strip(), pattern).date()
        except ValueError:
            pass
    raise ValueError('Data inválida')


def integer(value):
    try:
        result = Decimal(text(value))
    except InvalidOperation:
        raise ValueError('Duração ou presenças sem número válido') from None
    if not result.is_finite() or result != result.to_integral_value():
        raise ValueError('Número inteiro obrigatório')
    return int(result)


def source_json(headers, row):
    return {text(key): value.isoformat() if isinstance(value, (date, datetime)) else value
            for key, value in zip(headers, row) if key is not None}


class Command(BaseCommand):
    help = 'Importação Excel apenas por N.º Mecanográfico; simulação por defeito.'

    def add_arguments(self, parser):
        parser.add_argument('workbook', type=Path)
        parser.add_argument('--apply', action='store_true', help='Gravar os registos válidos após revisão da simulação.')
        parser.add_argument('--report', type=Path, required=True, help='Relatório JSON de contagens e linhas a rever.')

    def handle(self, *args, **options):
        source = Path(options['workbook']).resolve()
        report_path = Path(options['report']).resolve()
        if report_path == source or report_path.suffix.lower() != '.json':
            raise CommandError('O relatório deve ser um ficheiro .json separado do Excel original.')
        workbook = load_workbook(source, read_only=True, data_only=True)
        report = {'mode': 'apply' if options['apply'] else 'dry-run',
                  'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
                  'employees_created': 0, 'employees_existing': 0,
                  'trainings_created': 0, 'trainings_existing': 0, 'issues': []}
        try:
            if {'FUNCIONARIOS_EPAL','BASE_EPAL'} & set(workbook.sheetnames) or ('FUNCIONÁRIOS_ EPAL' in workbook.sheetnames and 'Formandos' not in workbook.sheetnames):
                from training.integrated_import import import_integrated
                report = import_integrated(workbook, options['apply'], source)
                report_path.parent.mkdir(parents=True, exist_ok=True)
                report_path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding='utf-8')
                self.stdout.write(json.dumps({k: len(v) if isinstance(v,list) else v for k,v in report.items()}, ensure_ascii=False))
                return
            if not {'FUNCIONÁRIOS_ EPAL', 'Formandos'}.issubset(workbook.sheetnames):
                raise CommandError('Folhas esperadas não encontradas.')
            with transaction.atomic():
                sheet = workbook['FUNCIONÁRIOS_ EPAL']
                headers = next(sheet.iter_rows(min_row=11, max_row=11, values_only=True))
                if headers[1] != 'N.º Mecan' or headers[3] != 'Nome':
                    raise CommandError('Estrutura do cadastro diferente da versão validada (linha 11).')
                seen = set()
                for index, row in enumerate(sheet.iter_rows(min_row=12, values_only=True), 12):
                    if not any(value is not None for value in row):
                        continue
                    try:
                        number, name = text(row[1]), text(row[3])
                        if not number or not name:
                            raise ValueError('Falta número mecanográfico ou nome')
                        if number in seen:
                            raise ValueError('Número mecanográfico repetido no cadastro de origem')
                        seen.add(number)
                        existing = Employee.objects.filter(number=number).first()
                        if existing and existing.name.strip().casefold() != name.casefold():
                            raise ValueError('Número já existente com nome diferente; revisão manual necessária')
                        if existing:
                            report['employees_existing'] += 1
                            continue
                        source_data = source_json(headers, row)
                        employee = Employee(number=number, name=name, birth_date=parse_date(row[4]),
                            phone=text(row[8]), admission_date=parse_date(row[9]), category=text(row[10]),
                            function=text(source_data.get('Função')), email=text(row[11]),
                            direction=text(row[12]), area=text(row[13]), source_data=source_data)
                        employee.full_clean()
                        employee.save()
                        report['employees_created'] += 1
                    except (ValueError, ValidationError, InvalidOperation) as error:
                        report['issues'].append({'sheet': sheet.title, 'row': index, 'reason': str(error)})
                sheet = workbook['Formandos']
                headers = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
                required = {'N.º Mecanográfico', 'Nome Formação', 'Data de Início', 'Data de Fim',
                            'Duração (dias)', 'Nº Presenças', 'Nº Hora por dia'}
                if not required.issubset(headers):
                    raise CommandError('Colunas obrigatórias de Formandos não encontradas.')
                occurrences = {}
                for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
                    if not any(value is not None for value in row):
                        continue
                    values = dict(zip(headers, row))
                    raw = source_json(headers, row)
                    digest = hashlib.sha256(json.dumps(raw, sort_keys=True, ensure_ascii=False).encode()).hexdigest()
                    occurrences[digest] = occurrences.get(digest, 0) + 1
                    source_key = f'{digest}:{occurrences[digest]}'
                    try:
                        if Training.objects.filter(source_key=source_key).exists():
                            report['trainings_existing'] += 1
                            continue
                        number = text(values.get('N.º Mecanográfico'))
                        employee = Employee.objects.filter(number=number).first()
                        if employee is None:
                            raise ValueError('Número mecanográfico sem correspondência válida no cadastro')
                        training = Training(employee=employee, title=text(values.get('Nome Formação')),
                            area=text(values.get('Área/Categoria')), trainer=text(values.get('Formador(a)')),
                            provider=text(values.get('Entidade Formadora')), location=text(values.get('Local de formação')),
                            start=parse_date(values.get('Data de Início')), end=parse_date(values.get('Data de Fim')),
                            days=integer(values.get('Duração (dias)')), attendance=integer(values.get('Nº Presenças')),
                            hours_per_day=Decimal(text(values.get('Nº Hora por dia'))),
                            grade=Decimal(text(values['Média Final'])) if values.get('Média Final') is not None else None,
                            certified=None,
                            status=text(values.get('Status-Curso')) or 'Concluído', source_key=source_key, source_data=raw)
                        training.full_clean()
                        training.save()
                        report['trainings_created'] += 1
                    except (ValueError, ValidationError, InvalidOperation) as error:
                        report['issues'].append({'sheet': sheet.title, 'row': index, 'reason': str(error)})
                if options['apply']:
                    AuditEvent.objects.create(actor='import_excel', action='IMPORT', detail=json.dumps(report, ensure_ascii=False))
                else:
                    transaction.set_rollback(True)
        finally:
            workbook.close()
        report_path.parent.mkdir(parents=True, exist_ok=True)
        report_path.write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding='utf-8')
        self.stdout.write(json.dumps({key: value if key != 'issues' else len(value)
                                     for key, value in report.items()}, ensure_ascii=False))
