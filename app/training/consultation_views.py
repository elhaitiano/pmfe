import io
from functools import wraps
from django.contrib.auth.decorators import login_required
from django.core.exceptions import PermissionDenied,ValidationError
from django.core.paginator import Paginator
from django.db import OperationalError
from django.http import FileResponse,HttpResponseBadRequest
from django.shortcuts import render,get_object_or_404,redirect
from django.views.decorators.http import require_http_methods
from django.views.decorators.cache import never_cache
from django.urls import reverse
from . import consultations as service
from . import evaluation_corrections as corrections
from .models import Training

def access(view):
    @login_required
    @never_cache
    @wraps(view)
    def guarded(request,*a,**kw):
        if not request.user.has_perms(['training.view_training','training.view_employee']):raise PermissionDenied
        return view(request,*a,**kw)
    return guarded
@access
@require_http_methods(['GET'])
def index(request):
    form=service.Filters(request.GET or None)
    valid=not form.is_bound or form.is_valid()
    rows=service.select(form.cleaned_data if form.is_bound and valid else {}) if valid else []
    totals=service.summary(rows);fmt=request.GET.get('format','')
    if fmt and not valid:return HttpResponseBadRequest('Corrija os filtros antes de exportar.')
    if fmt not in ('','xlsx','pdf'):return HttpResponseBadRequest('Formato desconhecido.')
    if fmt:return export(rows,totals,fmt,request.GET.urlencode())
    query=request.GET.copy();query.pop('page',None);query.pop('format',None)
    return render(request,'consultations/index.html',{'form':form,'page':Paginator(rows,40).get_page(request.GET.get('page')),'totals':totals,'query':query.urlencode(),'can_correct':request.user.is_staff and request.user.has_perm('training.change_training'),'can_documents':request.user.is_staff and request.user.has_perm('training.change_trainingdocumentstate')})
def export(rows,totals,fmt,filters):
    output=io.BytesIO()
    if fmt=='xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font,PatternFill
        from openpyxl.utils import get_column_letter
        wb=Workbook();ws=wb.active;ws.title='Participações';ws.append(service.HEADERS)
        for t in rows:
            values=service.export_row(t);ws.append(values)
            # Explicit string type prevents Excel formula injection from text.
            for cell,value in zip(ws[ws.max_row],values):
                if isinstance(value,str):cell.data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(color='FFFFFF',bold=True);cell.fill=PatternFill('solid',fgColor='103C62')
        for i in range(1,len(service.HEADERS)+1):ws.column_dimensions[get_column_letter(i)].width=25 if i!=4 else 65
        summary=wb.create_sheet('Resumo');summary.append(['Indicador','Valor'])
        for k,label in [('physical','Participações na lista'),('participations','Participações consolidadas'),('people','Trabalhadores únicos'),('numeric_people','Com avaliação numérica'),('qualitative_people','Com avaliação qualitativa'),('duplicates','Duplicados preservados')]:summary.append([label,totals[k]])
        summary.append(['Filtros',filters]);summary.cell(summary.max_row,2).data_type='s'
        for name,items in [('Direcções',totals['directions']),('Grupos',totals['groups']),('Formações',totals['formations']),('Anos',totals['years']),('Faixas etárias',totals['ages'])]:
            sheet=wb.create_sheet(name);sheet.append(['Identificação','Participações','Trabalhadores únicos','Notas numéricas','Média /20','Avaliações qualitativas'])
            for item in items:
                sheet.append([item[k] for k in ('label','participations','people','numeric_count','average','qualitative_count')]);sheet.cell(sheet.max_row,1).data_type='s'
        wb.save(output);wb.close();content='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        from reportlab.platypus import SimpleDocTemplate,Spacer
        from reportlab.lib.pagesizes import A4,landscape
        from .executive_pdf import para,table
        doc=SimpleDocTemplate(output,pagesize=landscape(A4),leftMargin=24,rightMargin=24,topMargin=24,bottomMargin=24)
        story=[para('EPAL / PROÁGUA — Consultas e Relatórios','ExecTitle'),para(f"{totals['people']} trabalhadores únicos; {totals['participations']} participações consolidadas; {totals['physical']} linhas detalhadas."),para('Filtros: '+filters,'ExecSmall'),Spacer(1,12)]
        data=[]
        for t in rows:
            data.append([t.employee.number+' — '+t.employee.name,t.title,t.start.isoformat()+' / '+t.end.isoformat(),str(t.completed_hours)+' H' if t.attendance_confirmed else 'Por confirmar',t.evaluation['label'],service.state(t.consultation_issued)+' / '+service.state(t.consultation_learner)+' / '+service.state(t.consultation_hr)])
        story.extend([para('Certificados: emissão / formando / RH. As notas calculadas usam percentagem × 20 / 100, apenas em avaliações finais comprovadas. Duplicados permanecem nas listas, excluídos dos totais.','ExecSmall'),table(['Formando','Formação','Período','Horas','Avaliação','Certificados'],data,widths=[125,180,105,65,140,175])])
        doc.build(story);content='application/pdf'
    output.seek(0);response=FileResponse(output,content_type=content,as_attachment=True,filename='PROAGUA-consultas.'+fmt)
    response['Cache-Control']='private, no-store';response['X-Content-Type-Options']='nosniff';return response

@access
@require_http_methods(['GET','POST'])
def correct(request,pk):
    corrections.require(request.user);training=get_object_or_404(Training,pk=pk)
    from .source_achievement import attach
    attach([training]);form=corrections.CorrectionForm(request.POST if request.method=='POST' and not request.POST.get('token') else None)
    from .source_achievement import index as source_index
    context={'training':training,'form':form,'original':source_index().get(training.source_key)}
    if request.method=='POST':
        try:
            if request.POST.get('token'):
                if request.POST.get('confirm')!='yes':raise ValidationError('Confirme expressamente a correcção.')
                event=corrections.confirm(request.user,request.POST['token'],training.pk)
                return redirect(reverse('consultations')+'?q='+training.employee.number)
            if form.is_valid():
                token,payload=corrections.preview(request.user,training,form.cleaned_data)
                from .evaluations import percentage_to_grade
                context.update(token=token,proposal=payload['proposed'],proposal_label=dict(corrections.CorrectionForm.base_fields['unit'].choices)[payload['proposed']['unit']],equivalent=percentage_to_grade(payload['proposed']['percentage']) if payload['proposed']['unit']=='percentage' else None)
        except (ValidationError,OperationalError) as exc:
            context['error']=str(exc) if isinstance(exc,ValidationError) else 'Operação concorrente. Reabra a participação.'
    return render(request,'consultations/correct.html',context)
